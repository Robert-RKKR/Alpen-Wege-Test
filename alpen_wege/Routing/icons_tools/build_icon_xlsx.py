#!/usr/bin/env python3
"""One-off: builds icon_set.xlsx from scripts/icon_sources/icon_map.csv and the generated icons_import.csv.
    python build_icon_xlsx.py ICON_MAP_CSV ICONS_IMPORT_CSV OUT_XLSX
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

map_csv, import_csv, out = sys.argv[1:4]
csv.field_size_limit(10_000_000)
icons = list(csv.DictReader(open(map_csv, encoding="utf-8")))
final = list(csv.DictReader(open(import_csv, encoding="utf-8")))
assert [r["key"] for r in icons] == [r["key"] for r in final]

F = "Arial"
head_font, head_fill = Font(name=F, bold=True, color="FFFFFF", size=10), PatternFill("solid", start_color="1F3A5F")
body = Font(name=F, size=10)
edit_fill = PatternFill("solid", start_color="FFF9E0")          # the columns you edit
thin = Side(style="thin", color="D9D9D9")
border = Border(bottom=thin)

wb = Workbook()

# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
rows = [
    ("Icons for Alpen Wege: one row per icon key (437 + 1)", True),
    ("", False),
    ("What is in this file", True),
    ("icons  =  the list to EDIT. One row per icon: names, and which graphic stands for it. Same columns as scripts/icon_sources/icon_map.csv.", False),
    ("import =  the FINISHED list with the SVG graphic of every icon (what scripts/make_icon_set.py makes from the icons sheet). Same columns as icons_import.csv.", False),
    ("", False),
    ("Import into the database (nothing to type into Excel)", True),
    ("1.  Sheet 'import': File > Save As > CSV UTF-8 (comma delimited), name it icons_import.csv.  (Or use the icons_import.csv that is next to this file.)", False),
    ("2.  BY API (nothing to copy to the host): POST /api/v1/assets/icons/upload-csv?dry_run=true with the file as form field 'file' (administrator token), then the same without dry_run. In Postman: assets / icons / Upload icons CSV (dry run) and (import all). With link=true (default) the rows get their icon_id in the same call.", False),
    ("2b. OR BY SCRIPT: copy the CSV to the host and run:   sudo docker exec alpen-api python /app/scripts/seed_icons.py --csv /app/icons_import.csv   (add --dry-run first)", False),
    ("3.  Link the reference tables (script way only; the API call links by itself):   sudo docker exec -i alpen-postgres psql -U alpenwege -d alpenwege < scripts/sql/icons_link.sql   or   POST /api/v1/assets/icons/link", False),
    ("4.  Check:   sudo docker exec -i alpen-postgres psql -U alpenwege -d alpenwege < scripts/sql/icons_list_keys.sql    (icon_exists = t and rows_linked = rows_using for every row)", False),
    ("", False),
    ("Change an icon", True),
    ("Edit the yellow columns in the sheet 'icons' (names, kind, value, color, tag), save that sheet as CSV UTF-8 over scripts/icon_sources/icon_map.csv,", False),
    ("then run:  python scripts/make_icon_set.py --out icons_import.csv --preview icon_preview.html   and import again with  seed_icons.py --csv icons_import.csv --update", False),
    ("", False),
    ("kind: what is drawn", True),
    ("lucide = a Lucide icon (value = its name, see lucide.dev; the file must be in scripts/icon_sources/lucide/). color empty = the screen colours it. tag = up to 3 letters in a corner label.", False),
    ("grade  = a rounded square with the text in value (T3, IV, WI4) in the colour of the row.     bars = value '3/7': 3 of 7 bars filled.     badge = a round badge with 1 to 3 letters.", False),
    ("Type value of a bars row as TEXT (the column is formatted as text): Excel would turn 3/7 into a date otherwise.", False),
    ("", False),
    ("quality (column K): 'exact Lucide name' = fine.  'mapped to Lucide' = my choice, check it.  'same glyph N times in <table>: redraw' = the same picture is used N times in that table; give them different pictures when you can.", False),
    ("Sources and licence: Lucide icons, ISC licence (scripts/icon_sources/lucide/LICENSE). Names in German and Polish come from the seed data (push_data) where one value uses the key; the others are empty.", False),
    ("The grade colours and the rarity colours are the colours of the same values in the seed data.", False),
]
for r, (text, bold) in enumerate(rows, start=1):
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(name=F, size=12 if r == 1 else 10, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 150

# counts with formulas (they recalculate when the icons sheet changes)
start = len(rows) + 2
ws.cell(row=start, column=1, value="Counts (formulas on the sheet 'icons')").font = Font(name=F, bold=True, size=10)
counts = [
    ("Icons (rows)", "=COUNTA(icons!A2:A1000)"),
    ("kind lucide", '=COUNTIF(icons!E2:E1000,"lucide")'),
    ("kind grade", '=COUNTIF(icons!E2:E1000,"grade")'),
    ("kind bars", '=COUNTIF(icons!E2:E1000,"bars")'),
    ("kind badge", '=COUNTIF(icons!E2:E1000,"badge")'),
    ("to redraw (same glyph several times in one table)", '=COUNTIF(icons!K2:K1000,"same glyph*")'),
    ("without a German name", "=COUNTA(icons!A2:A1000)-COUNTA(icons!C2:C1000)"),
]
for i, (label, formula) in enumerate(counts, start=start + 1):
    ws.cell(row=i, column=1, value=label).font = body
    ws.cell(row=i, column=2, value=formula).font = body
ws.column_dimensions["B"].width = 12

# ------------------------------------------------------------------ icons (edit)
ws2 = wb.create_sheet("icons")
cols = list(icons[0].keys())
widths = {"key": 28, "name": 30, "name_de": 28, "name_pl": 28, "kind": 9, "value": 22, "color": 10, "tag": 6, "used_by": 22, "rows_using": 10, "quality": 44}
text_cols = {"key", "value", "color", "tag"}
editable = {"name", "name_de", "name_pl", "kind", "value", "color", "tag"}
for j, name in enumerate(cols, start=1):
    c = ws2.cell(row=1, column=j, value=name)
    c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="left", vertical="center")
    ws2.column_dimensions[get_column_letter(j)].width = widths.get(name, 14)
for i, r in enumerate(icons, start=2):
    for j, name in enumerate(cols, start=1):
        v = r[name]
        c = ws2.cell(row=i, column=j, value=int(v) if name == "rows_using" else v)
        c.font, c.border = body, border
        if name in text_cols:
            c.number_format = "@"
        if name in editable:
            c.fill = edit_fill
        if name == "color" and v:
            c.fill = PatternFill("solid", start_color=v.lstrip("#").upper())
            lum = (int(v[1:3], 16) * 299 + int(v[3:5], 16) * 587 + int(v[5:7], 16) * 114) / 1000
            c.font = Font(name=F, size=10, color="000000" if lum > 150 else "FFFFFF")
last = len(icons) + 1
ws2.freeze_panes = "B2"
ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
dv = DataValidation(type="list", formula1='"lucide,grade,bars,badge"', allow_blank=False, showErrorMessage=True,
                    errorTitle="kind", error="kind is lucide, grade, bars or badge")
ws2.add_data_validation(dv)
dv.add(f"E2:E{last}")
ws2.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=['LEFT($K2,10)="same glyph"'], font=Font(name=F, size=10, color="9C5700"), fill=PatternFill("solid", bgColor="FFEB9C")))

# ------------------------------------------------------------------ import (final, with the SVG text)
ws3 = wb.create_sheet("import")
cols3 = ["key", "name", "name_de", "name_pl", "summary", "svg"]
for j, name in enumerate(cols3, start=1):
    c = ws3.cell(row=1, column=j, value=name)
    c.font, c.fill = head_font, head_fill
ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 28, 30
ws3.column_dimensions["C"].width = ws3.column_dimensions["D"].width = 24
ws3.column_dimensions["E"].width, ws3.column_dimensions["F"].width = 34, 90
longest = 0
for i, r in enumerate(final, start=2):
    for j, name in enumerate(cols3, start=1):
        c = ws3.cell(row=i, column=j, value=r[name])
        c.font, c.number_format = body, "@"
        c.alignment = Alignment(wrap_text=False, vertical="top")
    longest = max(longest, len(r["svg"]))
assert longest < 32000, longest          # the cell limit of Excel is 32,767 characters
ws3.freeze_panes = "B2"
ws3.auto_filter.ref = f"A1:F{len(final) + 1}"
wb.move_sheet("import", offset=0)
wb.save(out)
print("saved", out, "| rows", len(icons), "| longest svg", longest, "characters")
